"""Hand-check workbook: the 10-K statements line by line, next to the tool's pull.

    python -m creditmemo.handcheck make DVA          # writes handcheck/DVA_handcheck.xlsx
    python -m creditmemo.handcheck compare DVA       # reads it back, writes handcheck/DVA_result.md

Tabs:
  Summary            unit, target, instructions, match rate per tab and overall
  Income statement   each statement exactly as in the latest 10-K: same lines,
  Balance sheet      same order, the latest two years side by side:
  Cash flow          tool | 10-K (type here) | diff | OK
  Notes              items from the notes (lease cost, lease liabilities,
                     next-year maturities, cash taxes)
  Ratios             every ratio rebuilt in Excel from your typed figures,
                     using the same line mapping as the standard spread,
                     next to the tool's ratio

You type numbers exactly as the 10-K prints them, including parentheses as a
minus sign. The unit is read from the statement header ("$ in Thousands").
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from . import ROOT, SPREAD_DIR

HANDCHECK_DIR = ROOT / "handcheck"
TABS = [("income", "Income statement"), ("balance", "Balance sheet"), ("cashflow", "Cash flow")]
NOTES_TAB = "Notes"
NOTE_LABELS = {
    "operating_lease_cost": ("Operating lease cost", "Leases note: operating lease cost"),
    "operating_lease_liabilities": ("Operating lease liabilities, total", "Leases note: total operating lease liabilities"),
    "scheduled_principal": ("Debt due in the next 12 months", "Debt note: maturities table, first year"),
    "cash_taxes": ("Cash taxes paid", "Cash flow supplemental disclosure or tax note"),
}
FIRST_ROW = 6
COLS_PER_YEAR = 4
YEAR_START_COL = 3
UNIT_NAMES = {1e3: "thousands", 1e6: "millions", 1e9: "billions", 1.0: "dollars"}


def _path(ticker: str) -> Path:
    return HANDCHECK_DIR / f"{ticker}_handcheck.xlsx"


def _col(n: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def _std_line(spread, statement, key):
    return next(ln for ln in spread["standard"][statement] if ln["key"] == key)


def make(ticker: str, periods=None) -> Path:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    spread = json.loads((SPREAD_DIR / f"{ticker}.json").read_text())
    ar = spread["as_reported"]["statements"]
    all_labels = [p["label"] for p in spread["periods"]]
    labels = periods or all_labels[-2:]
    unit = ar["income"]["money_scale"]
    latest = spread["as_reported"]["filings"][0]

    F = "Arial"
    base, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    title = Font(name=F, size=14, bold=True)
    blue, green = Font(name=F, size=10, color="0000FF"), Font(name=F, size=10, color="008000")
    grey = Font(name=F, size=9, color="666666")
    link = Font(name=F, size=10, color="0563C1", underline="single")
    head_fill, year_fill = PatternFill("solid", fgColor="DDE3EA"), PatternFill("solid", fgColor="EEF1F5")
    yellow, red_fill = PatternFill("solid", fgColor="FFFF00"), PatternFill("solid", fgColor="F8CBAD")
    thin = Side(style="thin", color="B7B7B7")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    num, pct = '#,##0;(#,##0);"-"', '0.00%;(0.00%);0.00%'
    mult = '0.00"x";(0.00"x");"-"'

    wb = Workbook()
    sm = wb.active
    sm.title = "Summary"
    sm["A1"] = f"{spread['name']} ({ticker}): hand check against the 10-K"
    sm["A1"].font = title
    sm["A3"], sm["B3"] = "Unit you type in ($ per unit)", unit
    sm["C3"] = f"From the statement header: '$ in {UNIT_NAMES.get(unit, unit)}'"
    sm["A4"], sm["B4"] = "Match target", 0.01
    sm["A5"], sm["B5"] = "Years", ", ".join(labels)
    sm["A6"], sm["B6"] = "10-K", f"Filed {latest['filed']} (accession {latest['accn']})"
    sm["B3"].font, sm["B3"].fill, sm["B3"].number_format = blue, yellow, "#,##0"
    sm["B4"].font, sm["B4"].number_format = blue, "0.0%"
    sm["A8"] = "How to use"
    sm["B8"] = ("Each statement tab lists the statement exactly as the 10-K prints it. Open the 10-K at "
                "that statement and type every figure into the yellow cells as printed: a number in "
                "parentheses is negative, so (579,926) is typed as -579926. Leave a cell blank until "
                "you've checked it. Red = more than the target away from the tool's figure.")
    ex = next(r for r in ar["income"]["rows"] if not r["abstract"] and r["kind"] == "money")
    ex_val = ex["values"][all_labels.index(labels[-1])] / unit
    sm["A9"] = "Example"
    sm["B9"] = (f"Income statement tab, '{ex['label']}', {labels[-1]}: the 10-K prints {ex_val:,.0f}, "
                f"so type {ex_val:.0f}. Under 'Page / note' write e.g. 'Consolidated Statements of "
                f"Income, p. 72'.")
    for c in ("A3", "A4", "A5", "A6", "A8", "A9"):
        sm[c].font, sm[c].alignment = bold, Alignment(vertical="top")
    for c in ("B3", "B4", "B5", "B6"):
        sm[c].alignment = Alignment(horizontal="left", vertical="top")
    for c in ("B5", "B6", "C3"):
        sm[c].font = base
    for c in ("B8", "B9"):
        sm[c].font, sm[c].alignment = base, wrap
    sm.merge_cells("C3:F3")
    sm.merge_cells("B8:F8")
    sm.merge_cells("B9:F9")
    sm.row_dimensions[8].height = 58
    sm.row_dimensions[9].height = 32
    sm.column_dimensions["A"].width = 28
    for c in "BCDEF":
        sm.column_dimensions[c].width = 16
    for i, h in enumerate(["Tab", "Lines checked", "Within target", "Match rate"], 1):
        c = sm.cell(row=11, column=i, value=h)
        c.font, c.fill = bold, head_fill

    def year_headers(ws, row4, row5, extra):
        ws.cell(row=row5, column=1, value="Line")
        ws.cell(row=row5, column=2, value="XBRL tag")
        for yi, lab in enumerate(labels):
            c0 = YEAR_START_COL + yi * COLS_PER_YEAR
            ws.merge_cells(start_row=row4, start_column=c0, end_row=row4, end_column=c0 + COLS_PER_YEAR - 1)
            band = ws.cell(row=row4, column=c0, value=lab)
            band.font, band.fill, band.alignment = bold, year_fill, center
            for k, h in enumerate(["Tool", "10-K", "Diff", "OK"]):
                ws.cell(row=row5, column=c0 + k, value=h)
        tail = YEAR_START_COL + len(labels) * COLS_PER_YEAR
        for k, h in enumerate(extra):
            ws.cell(row=row5, column=tail + k, value=h)
        for cell in ws[row5]:
            if cell.value is not None:
                cell.font, cell.fill, cell.alignment = bold, head_fill, center
                cell.border = Border(bottom=thin)
        ws.freeze_panes = ws.cell(row=FIRST_ROW, column=YEAR_START_COL)
        return tail

    def value_cells(ws, r, raw_by_label, note=""):
        """Tool / 10-K / diff / OK for each year; returns {label: 10-K cell ref}."""
        refs = {}
        for yi, lab in enumerate(labels):
            c0 = YEAR_START_COL + yi * COLS_PER_YEAR
            T, K, D, O = (_col(c0 + k) for k in range(4))
            raw = raw_by_label.get(lab)
            t = ws.cell(row=r, column=c0, value=None if raw is None else f"={raw}/Summary!$B$3")
            t.font, t.number_format = blue, num
            if raw is not None:
                t.comment = Comment(f"Tool: {raw:,.0f} dollars. {note}".strip(), "creditmemo")
            k = ws.cell(row=r, column=c0 + 1)
            k.fill, k.font, k.number_format = yellow, blue, num
            d = ws.cell(row=r, column=c0 + 2,
                        value=f'=IF(OR({K}{r}="",{T}{r}=""),"",'
                              f'IF({K}{r}=0,IF({T}{r}=0,0,1),({T}{r}-{K}{r})/ABS({K}{r})))')
            d.font, d.number_format = base, pct
            o = ws.cell(row=r, column=c0 + 3, value=f'=IF({D}{r}="","",IF(ABS({D}{r})<=Summary!$B$4,1,0))')
            o.font, o.number_format, o.alignment = base, '"✓";;"✗"', center
            ws.conditional_formatting.add(
                f"{D}{r}:{O}{r}",
                FormulaRule(formula=[f'AND(${D}{r}<>"",ABS(${D}{r})>Summary!$B$4)'], fill=red_fill))
            refs[lab] = f"'{ws.title}'!${K}${r}"
        return refs

    def finish_tab(ws, tail, last, tab_i, name):
        widths = {1: 46, 2: 34}
        for yi in range(len(labels)):
            c0 = YEAR_START_COL + yi * COLS_PER_YEAR
            widths.update({c0: 13, c0 + 1: 13, c0 + 2: 9, c0 + 3: 5})
        widths.update({tail: 24, tail + 1: 22})
        for cnum, w in widths.items():
            ws.column_dimensions[_col(cnum)].width = w
        ws.column_dimensions[_col(tail + 2)].hidden = True
        dv = DataValidation(type="decimal", allow_blank=True)
        dv.error, dv.errorTitle = "Type a number, as printed in the 10-K", "Number expected"
        ws.add_data_validation(dv)
        k_cols = [_col(YEAR_START_COL + yi * COLS_PER_YEAR + 1) for yi in range(len(labels))]
        o_cols = [_col(YEAR_START_COL + yi * COLS_PER_YEAR + 3) for yi in range(len(labels))]
        for k in k_cols:
            dv.add(f"{k}{FIRST_ROW}:{k}{last}")
        sr = 12 + tab_i
        sm.cell(row=sr, column=1, value=name).font = base
        sm.cell(row=sr, column=2, value="=" + "+".join(
            f"COUNT('{name}'!{k}{FIRST_ROW}:{k}{last})" for k in k_cols)).font = green
        sm.cell(row=sr, column=3, value="=" + "+".join(
            f"SUM('{name}'!{o}{FIRST_ROW}:{o}{last})" for o in o_cols)).font = green
        c = sm.cell(row=sr, column=4, value=f'=IF(B{sr}=0,"",C{sr}/B{sr})')
        c.number_format, c.font = "0.0%", base
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    cell_of = {}   # (statement, row id, label) -> 10-K cell
    for tab_i, (kind, name) in enumerate(TABS):
        stmt = ar[kind]
        ws = wb.create_sheet(name)
        ws["A1"] = f"{stmt['short_name']} ({spread['name']})"
        ws["A1"].font = title
        ws["A2"] = (f"As printed in the 10-K filed {latest['filed']}: {stmt['title']}. "
                    "Yellow = type the 10-K figure. Blue = the tool's figure. Red = outside the target.")
        ws["A2"].font = grey
        url = stmt["periods"][-1]["url"]
        ws["A3"] = "Open this statement on SEC.gov"
        ws["A3"].hyperlink, ws["A3"].font = url, link
        tail = year_headers(ws, 4, 5, ["Page / note", "Note", "row id"])
        idx = {lab: i for i, lab in enumerate(p["label"] for p in stmt["periods"])}
        r = FIRST_ROW
        for row in stmt["rows"]:
            if row["abstract"]:
                h = ws.cell(row=r, column=1, value=row["label"])
                h.font = Font(name=F, size=10, bold=True, italic=True, color="44546A")
                r += 1
                continue
            if row.get("kind") != "money":
                continue
            a = ws.cell(row=r, column=1, value=row["label"])
            a.font, a.alignment = (bold if row.get("total") else base), Alignment(indent=0 if row.get("total") else 1)
            ws.cell(row=r, column=2, value=row["tag"]).font = grey
            raw = {lab: row["values"][idx[lab]] for lab in labels if lab in idx}
            refs = value_cells(ws, r, raw)
            for lab, ref in refs.items():
                cell_of[(kind, row["id"], lab)] = ref
            pg = ws.cell(row=r, column=tail)
            pg.fill, pg.font = yellow, blue
            if row.get("only_in_prior"):
                ws.cell(row=r, column=tail + 1, value="Line appears only in the prior 10-K").font = base
            ws.cell(row=r, column=tail + 2, value=f"{kind}:{row['id']}").font = grey
            for cc in range(1, tail + 2):
                ws.cell(row=r, column=cc).border = Border(bottom=Side(style="hair", color="D9D9D9"))
            r += 1
        finish_tab(ws, tail, r - 1, tab_i, name)

    # Notes tab
    ws = wb.create_sheet(NOTES_TAB)
    ws["A1"], ws["A1"].font = f"Items from the notes ({spread['name']})", title
    ws["A2"] = "These come from the XBRL company facts, not the statements. Type them from the notes."
    ws["A2"].font = grey
    tail = year_headers(ws, 4, 5, ["Page / note", "Tool note", "row id"])
    ws.cell(row=5, column=2, value="XBRL tag / source")
    note_cell = {}
    r = FIRST_ROW
    for key, (lab_txt, hint) in NOTE_LABELS.items():
        recs = spread["notes"][key]
        ws.cell(row=r, column=1, value=f"{lab_txt} ({hint})").font = base
        srcs = {(recs.get(l) or {}).get("tag") or (recs.get(l) or {}).get("source") for l in labels}
        ws.cell(row=r, column=2, value="; ".join(sorted(s for s in srcs if s))).font = grey
        raw = {l: (recs.get(l) or {}).get("value") for l in labels}
        refs = value_cells(ws, r, raw)
        for l, ref in refs.items():
            note_cell[(key, l)] = ref
        notes_txt = {(recs.get(l) or {}).get("note") for l in labels} - {None, ""}
        if any((recs.get(l) or {}).get("proxy") for l in labels):
            notes_txt.add("PROXY")
        pg = ws.cell(row=r, column=tail)
        pg.fill, pg.font = yellow, blue
        ws.cell(row=r, column=tail + 1, value=" ".join(sorted(notes_txt))).font = base
        ws.cell(row=r, column=tail + 2, value=f"notes:{key}").font = grey
        r += 1
    finish_tab(ws, tail, r - 1, len(TABS), NOTES_TAB)

    tr = 12 + len(TABS) + 1
    sm.cell(row=tr, column=1, value="All tabs").font = bold
    sm.cell(row=tr, column=2, value=f"=SUM(B12:B{tr - 1})").font = bold
    sm.cell(row=tr, column=3, value=f"=SUM(C12:C{tr - 1})").font = bold
    c = sm.cell(row=tr, column=4, value=f'=IF(B{tr}=0,"",C{tr}/B{tr})')
    c.font, c.number_format = bold, "0.0%"

    # Ratios tab: standard lines rebuilt from typed cells through the same mapping
    rs = wb.create_sheet("Ratios")
    rs["A1"], rs["A1"].font = "Ratios rebuilt from your typed 10-K figures", title
    rs["A2"] = ("Each line adds up the typed statement lines that the standard spread maps to it, with "
                "the same signs. A year fills in once its lines are typed. Green = linked cells.")
    rs["A2"].font = grey
    for yi, lab in enumerate(labels):
        c0 = 2 + yi * 3
        rs.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 2)
        band = rs.cell(row=4, column=c0, value=lab)
        band.font, band.fill, band.alignment = bold, year_fill, center
        for k, h in enumerate(["From 10-K", "Tool", "Diff"]):
            cc = rs.cell(row=5, column=c0 + k, value=h)
            cc.font, cc.fill, cc.alignment = bold, head_fill, center
    rs.cell(row=5, column=1, value="Item").font = bold
    rs.cell(row=5, column=1).fill = head_fill
    rs.column_dimensions["A"].width = 40
    for i in range(2, 2 + 3 * len(labels)):
        rs.column_dimensions[_col(i)].width = 13
    rs.freeze_panes = "B6"

    def mapped(statement, key, lab, extra_sign=1.0):
        """Formula pieces (cells) for a standard line in one year."""
        ln = _std_line(spread, statement, key)
        per = spread["standard"]["periods"][statement]
        p = per.index(lab) if lab in per else None
        parts = []
        for s in ln["sources"]:
            if s.get("periods") is not None and p not in s["periods"]:
                continue
            ref = cell_of.get((statement, s["id"], lab))
            if ref is None:
                continue
            sign = s["sign"] * extra_sign
            parts.append(("-" if sign < 0 else "+") + ref)
        return parts

    def single(statement, key, lab):
        ps = mapped(statement, key, lab)
        return ps[0].lstrip("+") if len(ps) == 1 and not ps[0].startswith("-") else None

    state = {"r": 6}
    built = {}

    def put(label, key, kind, fn, tool_fn):
        r = state["r"]
        rs.cell(row=r, column=1, value=label).font = bold if kind == "ratio" else base
        for yi, lab in enumerate(labels):
            c0 = 2 + yi * 3
            fx = rs.cell(row=r, column=c0, value=fn(lab))
            fx.font, fx.number_format = green, mult if kind == "ratio" else num
            tv = tool_fn(lab)
            t = rs.cell(row=r, column=c0 + 1, value=tv)
            t.font, t.number_format = blue, mult if kind == "ratio" else num
            a_, b_ = fx.coordinate, t.coordinate
            d = rs.cell(row=r, column=c0 + 2,
                        value=f'=IF(OR({a_}="",{b_}="",{a_}="n.m."),"",IF({a_}=0,"",({b_}-{a_})/ABS({a_})))')
            d.font, d.number_format = base, pct
        built[key] = r
        state["r"] += 1

    def cell(key, lab):
        return f"{_col(2 + 3 * labels.index(lab))}{built[key]}"

    def sum_formula(parts):
        refs = [p[1:] for p in parts]
        if not refs:
            return '=""'
        guard = ",".join(f'{x}=""' for x in refs)
        return f'=IF(OR({guard}),"",{"".join(parts).lstrip("+")})'

    tool_amt = lambda k: (lambda l: None if spread["derived"][l].get(k) is None
                          else f"={spread['derived'][l][k]}/Summary!$B$3")
    tool_in = lambda k: (lambda l: None if spread["ratio_inputs"][l].get(k) is None
                         else f"={spread['ratio_inputs'][l][k]}/Summary!$B$3")
    tool_ratio = lambda k: (lambda l: spread["ratios"][l][k])

    put("Revenue", "revenue", "amt", lambda l: sum_formula(mapped("income", "revenue", l)), tool_in("revenue"))
    put("Operating income", "oi", "amt", lambda l: sum_formula(mapped("income", "operating_income", l)),
        tool_in("operating_income"))
    put("D&A (cash flow statement)", "da", "amt", lambda l: sum_formula(mapped("cashflow", "da", l)),
        tool_in("da"))
    put("EBITDA", "ebitda", "amt",
        lambda l: f'=IF(OR({cell("oi", l)}="",{cell("da", l)}=""),"",{cell("oi", l)}+{cell("da", l)})',
        tool_amt("ebitda"))
    put("Interest expense", "interest", "amt",
        lambda l: sum_formula(mapped("income", "interest_expense", l)), tool_in("interest_expense"))
    put("Total debt", "debt", "amt",
        lambda l: sum_formula(mapped("balance", "debt_current", l) + mapped("balance", "ltd", l)),
        tool_amt("total_debt"))
    put("Cash", "cash", "amt", lambda l: sum_formula(mapped("balance", "cash", l)), tool_in("cash"))
    put("Net debt", "net", "amt",
        lambda l: f'=IF(OR({cell("debt", l)}="",{cell("cash", l)}=""),"",{cell("debt", l)}-{cell("cash", l)})',
        tool_amt("net_debt"))
    put("Capex", "capex", "amt", lambda l: sum_formula(mapped("cashflow", "capex", l, -1.0)), tool_in("capex"))
    for key, (lab_txt, _) in NOTE_LABELS.items():
        put(f"{lab_txt} (Notes tab)", key, "amt",
            (lambda k: lambda l: f'=IF({note_cell[(k, l)]}="","",{note_cell[(k, l)]})')(key), tool_in(key))
    put("EBITDAR", "ebitdar", "amt",
        lambda l: (f'=IF(OR({cell("ebitda", l)}="",{cell("operating_lease_cost", l)}=""),"",'
                   f'{cell("ebitda", l)}+{cell("operating_lease_cost", l)})'), tool_amt("ebitdar"))
    put("Current assets", "ca", "amt", lambda l: sum_formula(mapped("balance", "current_assets", l)),
        tool_in("current_assets"))
    put("Current liabilities", "cl", "amt", lambda l: sum_formula(mapped("balance", "current_liabilities", l)),
        tool_in("current_liabilities"))
    state["r"] += 1

    def ratio(numer, denom, needs):
        def fn(l):
            refs = [cell(k, l) for k in needs]
            guard = ",".join(f'{x}=""' for x in refs)
            return f'=IF(OR({guard}),"",IF({denom(l)}<=0,"n.m.",({numer(l)})/({denom(l)})))'
        return fn

    C = lambda k: (lambda l: cell(k, l))
    put("Debt / EBITDA", "r1", "ratio", ratio(C("debt"), C("ebitda"), ["debt", "ebitda"]),
        tool_ratio("debt_ebitda"))
    put("Net debt / EBITDA", "r2", "ratio", ratio(C("net"), C("ebitda"), ["net", "ebitda"]),
        tool_ratio("net_debt_ebitda"))
    put("(Debt + leases) / EBITDAR", "r3", "ratio",
        ratio(lambda l: f"{cell('debt', l)}+{cell('operating_lease_liabilities', l)}", C("ebitdar"),
              ["debt", "operating_lease_liabilities", "ebitdar"]), tool_ratio("lease_adj_debt_ebitdar"))
    put("EBITDA / interest", "r4", "ratio", ratio(C("ebitda"), C("interest"), ["ebitda", "interest"]),
        tool_ratio("interest_coverage"))
    put("DSCR", "r5", "ratio",
        ratio(lambda l: f"{cell('ebitda', l)}-{cell('capex', l)}-{cell('cash_taxes', l)}",
              lambda l: f"{cell('interest', l)}+{cell('scheduled_principal', l)}",
              ["ebitda", "capex", "cash_taxes", "interest", "scheduled_principal"]), tool_ratio("dscr"))
    put("Current ratio", "r6", "ratio", ratio(C("ca"), C("cl"), ["ca", "cl"]), tool_ratio("current_ratio"))
    rs.page_setup.orientation = "landscape"

    HANDCHECK_DIR.mkdir(parents=True, exist_ok=True)
    out = _path(ticker)
    wb.save(out)
    return out


def compare(ticker: str, target: float = 0.01) -> str:
    """Read typed 10-K figures back and write a markdown result for the README.

    Differences are recomputed in Python from the typed values, so Excel
    doesn't need to have recalculated the file.
    """
    from openpyxl import load_workbook

    wb = load_workbook(_path(ticker), data_only=True)  # typed values, or what typed formulas evaluate to
    spread = json.loads((SPREAD_DIR / f"{ticker}.json").read_text())
    ar = spread["as_reported"]["statements"]
    unit = wb["Summary"]["B3"].value
    rows = []
    for kind, name in TABS + [("notes", NOTES_TAB)]:
        ws = wb[name]
        labels, c = [], YEAR_START_COL
        while ws.cell(row=4, column=c).value:
            labels.append(ws.cell(row=4, column=c).value)
            c += COLS_PER_YEAR
        tail = YEAR_START_COL + len(labels) * COLS_PER_YEAR
        for r in range(FIRST_ROW, ws.max_row + 1):
            rid = ws.cell(row=r, column=tail + 2).value
            if not rid:
                continue
            page = ws.cell(row=r, column=tail).value
            line = ws.cell(row=r, column=1).value
            for yi, lab in enumerate(labels):
                typed = ws.cell(row=r, column=YEAR_START_COL + yi * COLS_PER_YEAR + 1).value
                if typed is None or typed == "":
                    continue
                st, key = rid.split(":", 1)
                if st == "notes":
                    raw = (spread["notes"][key].get(lab) or {}).get("value")
                else:
                    stmt = ar[st]
                    row = next(x for x in stmt["rows"] if x["id"] == int(key))
                    p = [x["label"] for x in stmt["periods"]].index(lab)
                    raw = row["values"][p]
                tool = None if raw is None else raw / unit
                typed = float(typed)
                if tool is None:
                    diff = None if typed != 0 else 0.0
                elif typed == 0:
                    diff = 0.0 if tool == 0 else 1.0
                else:
                    diff = (tool - typed) / abs(typed)
                ok = diff is not None and abs(diff) <= target
                rows.append((name, line, lab, tool, typed, diff, ok, page))

    n, hit = len(rows), sum(1 for x in rows if x[6])
    out = [f"## {ticker} hand check", "",
           f"**{hit} of {n} lines within {target:.0%} of the 10-K "
           f"({(hit / n if n else 0):.0%} match rate).** Values in $ {UNIT_NAMES.get(unit, unit)}.", ""]
    for _, name in TABS + [("notes", NOTES_TAB)]:
        sub = [x for x in rows if x[0] == name]
        if not sub:
            continue
        out += [f"### {name}", "", "| Line | Year | Tool | 10-K | Diff | Page / note |",
                "|---|---|---:|---:|---:|---|"]
        for _, line, year, t, k, d, ok, page in sub:
            ts = "missing" if t is None else f"{t:,.0f}"
            ds = "n/a" if d is None else f"{d:+.2%}"
            out.append(f"| {line} | {year} | {ts} | {k:,.0f} | {ds}{'' if ok else ' ⚠'} | {page or ''} |")
        out.append("")
    if hit < n:
        out.append("Lines marked ⚠ are outside the target and need an explanation.")
    text = "\n".join(out) + "\n"
    (HANDCHECK_DIR / f"{ticker}_result.md").write_text(text)
    return text


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("make")
    m.add_argument("ticker")
    m.add_argument("--periods", nargs="*", help="default: latest two fiscal years")
    c = sub.add_parser("compare")
    c.add_argument("ticker")
    args = ap.parse_args(argv)
    t = args.ticker.upper()
    print(make(t, args.periods) if args.cmd == "make" else compare(t))
    return 0


if __name__ == "__main__":
    sys.exit(main())
